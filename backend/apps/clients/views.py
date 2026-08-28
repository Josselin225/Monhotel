import csv
import io
from mon_hotel_backend.csv_utils import sanitize_csv_cell
from rest_framework import viewsets, filters, status, generics
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from apps.accounts.permissions import IsAdminOrManager
from rest_framework.response import Response
from django_filters.rest_framework import DjangoFilterBackend
from django.http import HttpResponse
from django.db.models import Sum, Count, Subquery, OuterRef, Prefetch
from django.db.models.functions import ExtractMonth, ExtractDay
from .models import Client, ClientNote
from .serializers import ClientSerializer, ClientNoteSerializer
from apps.tenants.mixins import HotelScopeMixin


class ClientViewSet(HotelScopeMixin, viewsets.ModelViewSet):
    serializer_class = ClientSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['nationality', 'id_type', 'vip_status']
    search_fields = ['first_name', 'last_name', 'email', 'phone', 'id_number']
    ordering_fields = ['last_name', 'first_name', 'created_at', 'vip_status']

    def get_permissions(self):
        if self.action in ('destroy', 'import_csv', 'toggle_blacklist'):
            return [IsAdminOrManager()]
        return super().get_permissions()

    def get_queryset(self):
        from apps.bookings.models import Booking
        bookings_prefetch = Prefetch(
            'bookings',
            queryset=Booking.objects.exclude(status='cancelled').only(
                'id', 'client_id', 'check_in', 'check_out', 'total_price', 'status'
            ),
        )
        qs = Client.objects.prefetch_related(bookings_prefetch)
        hotel = self.get_hotel()
        if hotel is not None:
            qs = qs.filter(hotel=hotel)

        last_stay_after  = self.request.query_params.get('last_stay_after')
        last_stay_before = self.request.query_params.get('last_stay_before')
        no_stays         = self.request.query_params.get('no_stays')

        if no_stays == 'true':
            qs = qs.filter(bookings__isnull=True)
            return qs

        if last_stay_after or last_stay_before:
            last_checkout = Booking.objects.filter(
                client=OuterRef('pk')
            ).exclude(status='cancelled').order_by('-check_out').values('check_out')[:1]
            qs = qs.annotate(last_stay_checkout=Subquery(last_checkout))
            if last_stay_after:
                qs = qs.filter(last_stay_checkout__gte=last_stay_after)
            if last_stay_before:
                qs = qs.filter(last_stay_checkout__lte=last_stay_before)

        return qs
    ordering = ['last_name']

    @action(detail=True, methods=['post'])
    def toggle_blacklist(self, request, pk=None):
        client = self.get_object()
        client.is_blacklisted   = not client.is_blacklisted
        client.blacklist_reason = request.data.get('reason', client.blacklist_reason)
        client.save(update_fields=['is_blacklisted', 'blacklist_reason'])
        return Response(ClientSerializer(client).data)

    @action(detail=False, methods=['post'])
    def import_csv(self, request):
        file = request.FILES.get('file')
        if not file:
            return Response({'detail': 'Fichier requis.'}, status=status.HTTP_400_BAD_REQUEST)
        if not file.name.lower().endswith('.csv'):
            return Response({'detail': 'Format CSV requis (.csv).'}, status=status.HTTP_400_BAD_REQUEST)
        if file.size > 5 * 1024 * 1024:
            return Response({'detail': 'Fichier trop volumineux (max 5 Mo).'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            decoded = file.read().decode('utf-8-sig')
        except UnicodeDecodeError:
            decoded = file.read().decode('latin-1')

        reader = csv.DictReader(decoded.splitlines(), delimiter=';')
        hotel = self.get_hotel()
        created = 0
        skipped = 0
        errors  = []
        from django.db import transaction as db_transaction
        with db_transaction.atomic():
            for i, row in enumerate(reader, start=2):
                if i - 1 > 2000:
                    errors.append('Import limité à 2000 lignes par fichier.')
                    break
                email = (row.get('Email') or row.get('email') or '').strip()
                phone = (row.get('Téléphone') or row.get('phone') or row.get('Tel') or '').strip()
                last  = (row.get('Nom') or row.get('last_name') or '').strip()
                first = (row.get('Prénom') or row.get('first_name') or '').strip()
                if not last:
                    errors.append(f'Ligne {i}: Nom manquant')
                    skipped += 1
                    continue
                if email and Client.objects.filter(email=email, hotel=hotel).exists():
                    skipped += 1
                    continue
                Client.objects.create(
                    hotel=hotel,
                    last_name=last,
                    first_name=first,
                    email=email,
                    phone=phone,
                    nationality=(row.get('Nationalité') or row.get('nationality') or '').strip(),
                    address=(row.get('Adresse') or row.get('address') or '').strip(),
                )
                created += 1
        return Response({'created': created, 'skipped': skipped, 'errors': errors})

    @action(detail=False, methods=['get'], permission_classes=[IsAdminOrManager])
    def export_csv(self, request):
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="clients.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Nom', 'Prénom', 'Email', 'Téléphone', 'Nationalité',
                         'Type pièce', 'N° pièce', 'Adresse', 'VIP', 'Réservations', 'Client depuis'])
        for c in self.get_queryset():
            writer.writerow([
                sanitize_csv_cell(c.last_name), sanitize_csv_cell(c.first_name),
                sanitize_csv_cell(c.email), sanitize_csv_cell(c.phone),
                sanitize_csv_cell(c.nationality), c.get_id_type_display(),
                sanitize_csv_cell(c.id_number), sanitize_csv_cell(c.address),
                c.get_vip_status_display(), c.bookings.count(),
                c.created_at.strftime('%d/%m/%Y'),
            ])
        return response

    @action(detail=False, methods=['get'], permission_classes=[IsAdminOrManager])
    def export_xlsx(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = 'Clients'

        HDR_FILL = PatternFill('solid', fgColor='C9973A')
        HDR_FONT = Font(bold=True, color='FFFFFF', size=11)
        ALT_FILL = PatternFill('solid', fgColor='FFF8EC')
        BORDER   = Border(
            left=Side(style='thin', color='E5E7EB'),
            right=Side(style='thin', color='E5E7EB'),
            top=Side(style='thin', color='E5E7EB'),
            bottom=Side(style='thin', color='E5E7EB'),
        )
        headers = ['Nom', 'Prénom', 'Email', 'Téléphone', 'Nationalité',
                   'Type pièce', 'N° pièce', 'Adresse', 'Statut VIP', 'Réservations', 'Client depuis']
        ws.append(headers)
        for col_i, _ in enumerate(headers, 1):
            cell = ws.cell(1, col_i)
            cell.fill = HDR_FILL
            cell.font = HDR_FONT
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = BORDER
        ws.row_dimensions[1].height = 22

        for row_i, c in enumerate(self.get_queryset(), 2):
            row = [
                sanitize_csv_cell(c.last_name), sanitize_csv_cell(c.first_name),
                sanitize_csv_cell(c.email or ''), sanitize_csv_cell(c.phone or ''),
                sanitize_csv_cell(c.nationality or ''), c.get_id_type_display(),
                sanitize_csv_cell(c.id_number or ''), sanitize_csv_cell(c.address or ''),
                c.get_vip_status_display(),
                c.bookings.count(), c.created_at.strftime('%d/%m/%Y'),
            ]
            ws.append(row)
            fill = ALT_FILL if row_i % 2 == 0 else None
            for col_i in range(1, len(row) + 1):
                cell = ws.cell(row_i, col_i)
                if fill: cell.fill = fill
                cell.border = BORDER
                cell.alignment = Alignment(vertical='center')

        col_widths = [20, 16, 30, 16, 16, 14, 16, 28, 12, 12, 14]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        ws.freeze_panes = 'A2'
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        resp = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        resp['Content-Disposition'] = 'attachment; filename="clients.xlsx"'
        return resp

    @action(detail=False, methods=['get'])
    def stats(self, request):
        from django.core.cache import cache
        from django.utils import timezone
        from datetime import timedelta

        def _int_param(name, default, lo, hi):
            try:
                return max(lo, min(hi, int(request.query_params.get(name, default))))
            except (TypeError, ValueError):
                return default

        top_limit       = _int_param('top_limit', 10, 1, 50)
        birthdays_limit = _int_param('birthdays_limit', 5, 1, 50)
        birthdays_days  = _int_param('birthdays_days', 30, 1, 365)

        now = timezone.now()
        hotel = self.get_hotel()
        cache_key = f'client_stats_{hotel.pk if hotel else "none"}_{now.strftime("%Y-%m-%d-%H")}_{top_limit}_{birthdays_limit}_{birthdays_days}'
        cached = cache.get(cache_key)
        if cached:
            return Response(cached)
        month_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)

        qs = Client.objects.filter(hotel=hotel) if hotel is not None else Client.objects.all()
        total = qs.count()
        vip_count = qs.filter(vip_status='vip').count()
        vvip_count = qs.filter(vip_status='vvip').count()
        new_this_month = qs.filter(created_at__gte=month_start).count()

        # Top clients by lifetime value — single query with dual annotation
        top_clients = []
        for c in (
            qs
            .annotate(lv=Sum('bookings__total_price'), bk_count=Count('bookings'))
            .order_by('-lv')[:top_limit]
        ):
            top_clients.append({
                'id': c.id,
                'full_name': c.full_name,
                'vip_status': c.vip_status,
                'lifetime_value': float(c.lv or 0),
                'bookings_count': c.bk_count or 0,
            })

        # Anniversaires à venir (fenêtre configurable) — filtrage BDD d'abord, calcul Python ensuite
        today = now.date()
        months = {today.month}
        nxt = today.replace(day=28)  # safe way to advance month
        import datetime
        months_ahead = birthdays_days // 28 + 2
        for _ in range(months_ahead):
            nxt = (nxt + datetime.timedelta(days=4)).replace(day=1)
            months.add(nxt.month)

        birthday_qs = (
            qs
            .exclude(birthday__isnull=True)
            .annotate(bday_month=ExtractMonth('birthday'), bday_day=ExtractDay('birthday'))
            .filter(bday_month__in=months)
            .only('id', 'first_name', 'last_name', 'birthday')
        )
        birthdays = []
        for c in birthday_qs:
            try:
                bday = c.birthday.replace(year=today.year)
                if bday < today:
                    bday = bday.replace(year=today.year + 1)
                days_until = (bday - today).days
                if 0 <= days_until <= birthdays_days:
                    birthdays.append({
                        'id': c.id,
                        'full_name': c.full_name,
                        'birthday': c.birthday.isoformat(),
                        'days_until': days_until,
                    })
            except ValueError:
                pass
        birthdays.sort(key=lambda x: x['days_until'])

        result = {
            'total': total,
            'vip_count': vip_count,
            'vvip_count': vvip_count,
            'new_this_month': new_this_month,
            'top_clients': top_clients,
            'upcoming_birthdays': birthdays[:birthdays_limit],
        }
        cache.set(cache_key, result, timeout=300)
        return Response(result)


class ClientNoteListCreateView(generics.ListCreateAPIView):
    serializer_class = ClientNoteSerializer
    permission_classes = [IsAuthenticated]

    def _client_queryset(self):
        hotel = getattr(self.request.user, 'hotel', None)
        qs = Client.objects.all()
        if hotel is not None:
            qs = qs.filter(hotel=hotel)
        return qs

    def get_queryset(self):
        # 404 si le client n'existe pas / n'appartient pas à l'hôtel de l'appelant.
        client = generics.get_object_or_404(self._client_queryset(), pk=self.kwargs['client_pk'])
        return ClientNote.objects.filter(client=client).select_related('author')

    def perform_create(self, serializer):
        client = generics.get_object_or_404(self._client_queryset(), pk=self.kwargs['client_pk'])
        serializer.save(
            client=client,
            author=self.request.user,
        )


class ClientNoteDeleteView(generics.DestroyAPIView):
    serializer_class = ClientNoteSerializer
    permission_classes = [IsAdminOrManager]

    def get_queryset(self):
        hotel = getattr(self.request.user, 'hotel', None)
        qs = ClientNote.objects.filter(client_id=self.kwargs['client_pk'])
        if hotel is not None:
            qs = qs.filter(client__hotel=hotel)
        return qs
