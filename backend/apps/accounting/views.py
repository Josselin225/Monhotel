import csv
import io
from mon_hotel_backend.csv_utils import sanitize_csv_cell
from decimal import Decimal
from datetime import date, timedelta
from dateutil.relativedelta import relativedelta

from django.db.models import Sum, Q
from django.http import HttpResponse
from django.utils import timezone

from rest_framework import viewsets, filters
from rest_framework.decorators import action
from rest_framework.response import Response
from django_filters.rest_framework import DjangoFilterBackend

from apps.accounts.permissions import IsAdminOrManager
from apps.tenants.mixins import HotelScopeMixin
from .models import Transaction, Budget
from .serializers import TransactionSerializer, BudgetSerializer


def _zero(val):
    return val if val is not None else Decimal('0')


class TransactionViewSet(HotelScopeMixin, viewsets.ModelViewSet):
    # Toute la comptabilité (lecture incluse) est réservée à admin/manager —
    # c'est un registre financier complet (dépenses, recettes, moyens de paiement).
    queryset = Transaction.objects.select_related('booking').all()
    serializer_class = TransactionSerializer
    permission_classes = [IsAdminOrManager]
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['type', 'category', 'payment_method']
    search_fields    = ['description', 'reference', 'notes']
    ordering_fields  = ['date', 'amount', 'created_at']
    ordering         = ['-date', '-created_at']

    def get_queryset(self):
        qs = super().get_queryset()
        date_from = self.request.query_params.get('date_from')
        date_to   = self.request.query_params.get('date_to')
        if date_from:
            qs = qs.filter(date__gte=date_from)
        if date_to:
            qs = qs.filter(date__lte=date_to)
        return qs

    @action(detail=False, methods=['get'])
    def summary(self, request):
        """Totaux sur une période donnée."""
        qs = self.get_queryset()

        # Appliquer filtres de date via query params
        date_from = request.query_params.get('date_from')
        date_to   = request.query_params.get('date_to')
        period    = request.query_params.get('period', 'month')

        today = date.today()
        if not date_from and not date_to:
            if period == 'month':
                date_from = today.replace(day=1).isoformat()
                date_to   = today.isoformat()
            elif period == '3months':
                date_from = (today - relativedelta(months=3)).isoformat()
                date_to   = today.isoformat()
            elif period == 'year':
                date_from = today.replace(month=1, day=1).isoformat()
                date_to   = today.isoformat()

        if date_from:
            qs = qs.filter(date__gte=date_from)
        if date_to:
            qs = qs.filter(date__lte=date_to)

        income  = _zero(qs.filter(type='income').aggregate(s=Sum('amount'))['s'])
        expense = _zero(qs.filter(type='expense').aggregate(s=Sum('amount'))['s'])
        net     = income - expense

        # Breakdown par catégorie
        by_category = {}
        for t in qs.values('category', 'type').annotate(total=Sum('amount')):
            by_category[t['category']] = {
                'total': str(t['total']),
                'type':  t['type'],
            }

        return Response({
            'income':      str(income),
            'expense':     str(expense),
            'net':         str(net),
            'by_category': by_category,
            'date_from':   date_from,
            'date_to':     date_to,
        })

    @action(detail=False, methods=['get'])
    def monthly(self, request):
        """Recettes et dépenses des 12 derniers mois."""
        today = date.today()
        months = []
        hotel = self.get_hotel()
        base_qs = Transaction.objects.filter(hotel=hotel) if hotel is not None else Transaction.objects.all()
        for i in range(11, -1, -1):
            start = (today - relativedelta(months=i)).replace(day=1)
            end   = (start + relativedelta(months=1)) - timedelta(days=1)
            qs    = base_qs.filter(date__gte=start, date__lte=end)
            income  = _zero(qs.filter(type='income').aggregate(s=Sum('amount'))['s'])
            expense = _zero(qs.filter(type='expense').aggregate(s=Sum('amount'))['s'])
            months.append({
                'month':   start.strftime('%b %Y'),
                'month_short': start.strftime('%b'),
                'year':    start.year,
                'income':  str(income),
                'expense': str(expense),
                'net':     str(income - expense),
            })
        return Response(months)

    @action(detail=False, methods=['get'])
    def export_csv(self, request):
        """Export CSV des transactions filtrées."""
        qs = self.filter_queryset(self.get_queryset())
        response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
        response['Content-Disposition'] = 'attachment; filename="transactions.csv"'

        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Référence', 'Type', 'Catégorie', 'Date', 'Description',
                         'Montant (FCFA)', 'Mode de paiement', 'Réservation', 'Notes'])
        for t in qs:
            writer.writerow([
                t.reference, t.get_type_display(), t.get_category_display(),
                t.date, sanitize_csv_cell(t.description), t.amount,
                t.get_payment_method_display(),
                t.booking.reference if t.booking_id else '',
                sanitize_csv_cell(t.notes),
            ])
        return response

    @action(detail=False, methods=['get'])
    def export_xlsx(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        qs = self.filter_queryset(self.get_queryset())
        wb = Workbook()
        ws = wb.active
        ws.title = 'Transactions'

        HDR_FILL  = PatternFill('solid', fgColor='C9973A')
        HDR_FONT  = Font(bold=True, color='FFFFFF', size=11)
        INC_FILL  = PatternFill('solid', fgColor='ECFDF5')
        EXP_FILL  = PatternFill('solid', fgColor='FEF2F2')
        BORDER    = Border(
            left=Side(style='thin', color='E5E7EB'),
            right=Side(style='thin', color='E5E7EB'),
            top=Side(style='thin', color='E5E7EB'),
            bottom=Side(style='thin', color='E5E7EB'),
        )
        headers = ['Référence', 'Type', 'Catégorie', 'Date', 'Description',
                   'Montant (FCFA)', 'Mode de paiement', 'Réservation', 'Notes']
        ws.append(headers)
        for col_i, _ in enumerate(headers, 1):
            cell = ws.cell(1, col_i)
            cell.fill = HDR_FILL; cell.font = HDR_FONT
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = BORDER
        ws.row_dimensions[1].height = 22

        for row_i, t in enumerate(qs, 2):
            row = [
                t.reference, t.get_type_display(), t.get_category_display(),
                t.date, sanitize_csv_cell(t.description or ''),
                float(t.amount), t.get_payment_method_display(),
                t.booking.reference if t.booking_id else '',
                sanitize_csv_cell(t.notes or ''),
            ]
            ws.append(row)
            fill = INC_FILL if t.type == 'income' else EXP_FILL
            for col_i in range(1, len(row) + 1):
                cell = ws.cell(row_i, col_i)
                cell.fill = fill
                cell.border = BORDER
                cell.alignment = Alignment(vertical='center')

        col_widths = [16, 12, 20, 12, 32, 16, 18, 14, 28]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'

        buf = io.BytesIO()
        wb.save(buf); buf.seek(0)
        resp = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        resp['Content-Disposition'] = 'attachment; filename="transactions.xlsx"'
        return resp

    @action(detail=False, methods=['post'])
    def import_csv(self, request):
        file = request.FILES.get('file')
        if not file:
            return Response({'detail': 'Fichier requis.'}, status=400)
        if not file.name.lower().endswith('.csv'):
            return Response({'detail': 'Format CSV requis (.csv).'}, status=400)
        if file.size > 5 * 1024 * 1024:
            return Response({'detail': 'Fichier trop volumineux (max 5 Mo).'}, status=400)
        try:
            decoded = file.read().decode('utf-8-sig')
        except UnicodeDecodeError:
            decoded = file.read().decode('latin-1')

        reader = csv.DictReader(decoded.splitlines(), delimiter=';')
        TYPE_MAP     = {'recette': 'income', 'income': 'income', 'revenue': 'income',
                        'dépense': 'expense', 'expense': 'expense', 'depense': 'expense'}
        CATEGORY_MAP = {v.lower(): k for k, v in Transaction.Category.choices}

        hotel = self.get_hotel()
        created = 0
        skipped = 0
        errors  = []
        from django.db import transaction as db_transaction
        for i, row in enumerate(reader, start=2):
            if i - 1 > 2000:
                errors.append('Import limité à 2000 lignes par fichier.')
                break
            try:
                raw_type = (row.get('Type') or '').strip().lower()
                t_type   = TYPE_MAP.get(raw_type)
                if not t_type:
                    errors.append(f'Ligne {i}: type invalide "{raw_type}"'); skipped += 1; continue
                amount_str = (row.get('Montant (FCFA)') or row.get('Montant') or '0').replace(' ', '').replace(',', '.')
                amount = Decimal(amount_str)
                raw_cat  = (row.get('Catégorie') or row.get('Categorie') or '').strip().lower()
                category = CATEGORY_MAP.get(raw_cat, 'other_income' if t_type == 'income' else 'other_expense')
                dt_str   = (row.get('Date') or '').strip()
                try:
                    from datetime import datetime
                    dt = datetime.strptime(dt_str, '%d/%m/%Y').date() if '/' in dt_str else date.fromisoformat(dt_str)
                except ValueError:
                    dt = date.today()
                with db_transaction.atomic():
                    Transaction.objects.create(
                        hotel=hotel,
                        type=t_type,
                        category=category,
                        amount=amount,
                        date=dt,
                        description=(row.get('Description') or '').strip(),
                        payment_method=(row.get('Mode de paiement') or 'other').strip().lower() or 'other',
                        notes=(row.get('Notes') or '').strip(),
                    )
                created += 1
            except Exception as exc:
                errors.append(f'Ligne {i}: {exc}')
                skipped += 1
        return Response({'created': created, 'skipped': skipped, 'errors': errors})


class BudgetViewSet(HotelScopeMixin, viewsets.ModelViewSet):
    # Budgets = données financières sensibles, réservées à admin/manager.
    queryset = Budget.objects.all()
    serializer_class = BudgetSerializer
    permission_classes = [IsAdminOrManager]
    filterset_fields = ['category', 'period', 'year', 'month']
    ordering = ['year', 'month', 'category']

    @action(detail=False, methods=['get'])
    def with_actuals(self, request):
        """
        Returns budgets with actual spending and alert status.
        ?year=YYYY&month=MM
        """
        year  = int(request.query_params.get('year',  timezone.now().year))
        month = request.query_params.get('month')
        month = int(month) if month else timezone.now().month
        hotel = self.get_hotel()

        budgets = self.get_queryset().filter(year=year, month=month)
        result  = []
        for b in budgets:
            spent_qs = Transaction.objects.filter(
                category=b.category,
                type=Transaction.Type.EXPENSE,
                date__year=year,
                date__month=month,
            )
            if hotel is not None:
                spent_qs = spent_qs.filter(hotel=hotel)
            spent = spent_qs.aggregate(total=Sum('amount'))['total'] or Decimal('0')

            pct     = float(spent / b.amount * 100) if b.amount else 0
            alert   = pct >= b.alert_pct
            over    = pct >= 100
            result.append({
                **BudgetSerializer(b).data,
                'spent':   float(spent),
                'pct':     round(pct, 1),
                'alert':   alert,
                'over':    over,
                'remaining': float(b.amount - spent),
            })
        return Response(result)
